#!/usr/bin/env python3
"""Build Aura Reset Center stationery.

  output/stationery/Aura-Letterhead-Template.docx (+ .pdf)
  output/stationery/Aura-Business-Cards.pdf        (90 x 55 mm, front + back per person)
  output/stationery/Aura-Compliment-Slip.pdf       (DL 210 x 99 mm)
  output/stationery/Aura-Envelope-DL.pdf           (220 x 110 mm)
  output/stationery/Aura-Email-Signatures.html     (paste into Gmail / Outlook)
  output/stationery/previews/*.png
"""
import subprocess
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_documents import (BRONZE, COCOA, DISPLAY, ESPRESSO, TAUPE, add_inline, set_font,
                             setup_document, spacing)

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "stationery"
BRAND = ROOT / "brand"

PEOPLE = [
    ("Eutemia Stockman", "Founder & Chief Executive Officer"),
    ("Gurpreet Singh", "Chief Executive Officer"),
    ("Sahill Chandan", "General Manager"),
]
ADDRESS = "332 Jumeirah First, Jumeirah 1, Dubai, UAE"


def letterhead():
    doc = setup_document("")
    doc.core_properties.title = "The Aura Reset Center — Letterhead"
    body = [
        ("[Date]", TAUPE),
        ("", None),
        ("[Recipient name]", COCOA),
        ("[Title, Company]", COCOA),
        ("[Address line 1]", COCOA),
        ("[City, Country]", COCOA),
        ("", None),
    ]
    for text, color in body:
        p = doc.add_paragraph()
        if text:
            add_inline(p, text, color=color)
        spacing(p, after=0, line=1.3)
    p = doc.add_paragraph()
    r = p.add_run("Subject: ")
    set_font(r, size=10, color=ESPRESSO, bold=True)
    add_inline(p, "[Subject of letter]", color=ESPRESSO, bold=True)
    spacing(p, before=6, after=12)
    for text in ("Dear [Name],",
                 "[Body of letter. Type here in Jost 10 pt. The header and footer carry the logo, address and "
                 "legal details automatically on every page.]",
                 "[Second paragraph.]",
                 "With warm regards,"):
        p = doc.add_paragraph()
        add_inline(p, text)
        spacing(p, after=8, line=1.35)
    p = doc.add_paragraph()
    spacing(p, before=28, after=0)
    set_font(p.add_run("[Signature]"), "Caveat", 20, BRONZE)
    p = doc.add_paragraph()
    set_font(p.add_run("[Name]"), DISPLAY, 12, ESPRESSO)
    spacing(p, after=0)
    p = doc.add_paragraph()
    add_inline(p, "[Title — e.g. General Manager]", size=8.5, color=TAUPE)
    p = doc.add_paragraph()
    spacing(p, before=24)
    add_inline(p, "Signatories: Eutemia Stockman, Founder & Chief Executive Officer · Gurpreet Singh, "
                  "Chief Executive Officer · Sahill Chandan, General Manager. Delete this line before sending.",
               size=7.5, color=TAUPE, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    out = OUT / "Aura-Letterhead-Template.docx"
    doc.save(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(out)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


FONTS_CSS = f"""
@font-face{{font-family:Marcellus;src:url({(BRAND/'fonts/Marcellus-Regular.ttf').as_uri()})}}
@font-face{{font-family:Jost;font-weight:300;src:url({(BRAND/'fonts/Jost-Light.ttf').as_uri()})}}
@font-face{{font-family:Jost;font-weight:400;src:url({(BRAND/'fonts/Jost-Regular.ttf').as_uri()})}}
@font-face{{font-family:Jost;font-weight:500;src:url({(BRAND/'fonts/Jost-Medium.ttf').as_uri()})}}
*{{box-sizing:border-box;margin:0}}
body{{font-family:Jost,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.pg{{position:relative;overflow:hidden;page-break-after:always}}
.ring{{position:absolute;border-radius:50%;border:1px solid rgba(240,169,104,.55);
  box-shadow:0 0 18px 2px rgba(240,169,104,.35),inset 0 0 18px 2px rgba(255,211,160,.15)}}
"""


def cards_html():
    dark = (BRAND / "logo-dark.png").as_uri()
    light = (BRAND / "logo-light.png").as_uri()
    pages = []
    for name, title in PEOPLE:
        pages.append(f"""
<div class="pg front"><div class="ring" style="width:62mm;height:62mm;left:14mm;top:-3.5mm"></div>
  <img src="{dark}" style="position:absolute;width:52mm;left:19mm;top:15mm"></div>
<div class="pg back">
  <img src="{light}" style="position:absolute;width:22mm;right:5mm;top:5mm">
  <div style="position:absolute;left:7mm;top:9mm">
    <div style="font-family:Marcellus;font-size:13pt;color:#221913">{name}</div>
    <div style="font-size:6pt;letter-spacing:.22em;text-transform:uppercase;color:#A6754C;margin-top:1.2mm">{title}</div>
  </div>
  <div style="position:absolute;left:7mm;bottom:7mm;font-size:6.4pt;line-height:1.65;color:#4A3A2C;font-weight:400">
    <div><span style="color:#A6754C">T</span>&nbsp; [Phone] &nbsp;·&nbsp; <span style="color:#A6754C">W</span>&nbsp; [WhatsApp]</div>
    <div><span style="color:#A6754C">E</span>&nbsp; [email]@[domain]</div>
    <div>{ADDRESS}</div>
    <div style="color:#7A6753">[Website] &nbsp;·&nbsp; [Instagram]</div>
  </div>
  <div style="position:absolute;left:7mm;right:7mm;bottom:4mm;height:.3mm;background:#DCC6B0"></div>
</div>""")
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{FONTS_CSS}
@page{{size:90mm 55mm;margin:0}}
.pg{{width:90mm;height:55mm}}
.front{{background:radial-gradient(70% 90% at 50% 45%,#3a2a1e 0%,#221913 70%)}}
.back{{background:#FBF5EE}}
</style></head><body>{''.join(pages)}</body></html>"""


def slip_html():
    light = (BRAND / "logo-light.png").as_uri()
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{FONTS_CSS}
@page{{size:210mm 99mm;margin:0}}
.pg{{width:210mm;height:99mm;background:#FBF5EE}}
</style></head><body><div class="pg">
<div class="ring" style="width:120mm;height:120mm;right:-40mm;top:-10mm;opacity:.5"></div>
<img src="{light}" style="position:absolute;width:42mm;left:14mm;top:12mm">
<div style="position:absolute;left:16mm;top:52mm;font-family:Marcellus;font-size:20pt;color:#221913">With compliments</div>
<div style="position:absolute;left:16mm;top:64mm;font-size:8pt;color:#7A6753;letter-spacing:.3em;text-transform:uppercase">Reset · Heal · Elevate</div>
<div style="position:absolute;left:16mm;right:16mm;bottom:9mm;border-top:.3mm solid #DCC6B0;padding-top:2.5mm;
  font-size:6.5pt;color:#7A6753;display:flex;justify-content:space-between">
  <span>The Aura Reset Center · {ADDRESS}</span><span>[Phone] · [Email] · [Website]</span>
  <span>[Registered Company Name] · TL [___] · TRN [___]</span></div>
</div></body></html>"""


def envelope_html():
    light = (BRAND / "logo-light.png").as_uri()
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{FONTS_CSS}
@page{{size:220mm 110mm;margin:0}}
.pg{{width:220mm;height:110mm;background:#ffffff}}
</style></head><body><div class="pg">
<img src="{light}" style="position:absolute;width:30mm;left:10mm;top:8mm">
<div style="position:absolute;left:12mm;top:33mm;font-size:6.5pt;line-height:1.6;color:#7A6753">
  The Aura Reset Center<br>{ADDRESS}<br>[Website]</div>
<div style="position:absolute;left:0;right:0;bottom:0;height:3mm;background:linear-gradient(90deg,#C79A6D,#A6754C)"></div>
</div></body></html>"""


def signatures_html():
    rows = []
    for name, title in PEOPLE:
        rows.append(f"""
<p style="font-family:Arial,sans-serif;font-size:12px;color:#7A6753;margin:28px 0 6px">Signature — {name}</p>
<table cellpadding="0" cellspacing="0" border="0" style="font-family:'Jost',Arial,Helvetica,sans-serif;color:#4A3A2C">
 <tr><td style="padding-right:16px;border-right:1px solid #DCC6B0;vertical-align:middle">
   <img src="[LOGO_URL]" alt="The Aura Reset Center" width="120" style="display:block;width:120px"></td>
 <td style="padding-left:16px;vertical-align:middle">
   <div style="font-family:'Marcellus',Georgia,serif;font-size:17px;color:#221913">{name}</div>
   <div style="font-size:10px;letter-spacing:2px;text-transform:uppercase;color:#A6754C;padding:3px 0 8px">{title}</div>
   <div style="font-size:12px;line-height:18px">T [Phone] &nbsp;|&nbsp; WhatsApp [WhatsApp]<br>
   <a href="mailto:[email]" style="color:#4A3A2C;text-decoration:none">[email]@[domain]</a> &nbsp;|&nbsp;
   <a href="[Website]" style="color:#A6754C;text-decoration:none">[Website]</a><br>
   <span style="color:#7A6753">{ADDRESS}</span></div>
 </td></tr>
 <tr><td colspan="2" style="padding-top:10px;font-size:9px;line-height:13px;color:#7A6753;max-width:520px">
   Reset · Heal · Elevate &nbsp;—&nbsp; Opening November 2026 in Jumeirah 1. This email and any attachments are
   confidential and intended only for the addressee. If you received it in error, please delete it and let us know.</td></tr>
</table>""")
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>Aura email signatures</title>
<style>body{{background:#ffffff;margin:24px;font-family:Arial,sans-serif}}</style></head><body>
<h2 style="font-family:Georgia,serif;color:#221913;font-weight:400">The Aura Reset Center — email signatures</h2>
<p style="font-size:13px;color:#4A3A2C;max-width:640px">How to use: open this file in a browser, select one signature
block, copy and paste it into Gmail (Settings → See all settings → Signature) or Outlook (Settings → Mail → Compose and reply).
Replace [LOGO_URL] with the logo hosted on the company website, and fill in the placeholders in square brackets.</p>
{''.join(rows)}</body></html>"""


RENDER_JS = r"""
const {chromium} = require('playwright');
(async () => {
  const [out, ...names] = process.argv.slice(1);
  const b = await chromium.launch();
  for (const n of names) {
    const p = await b.newPage({deviceScaleFactor: 2});
    await p.goto('file://' + out + '/.' + n + '.html');
    await p.evaluate(() => document.fonts.ready);
    await p.pdf({path: out + '/' + n + '.pdf', preferCSSPageSize: true, printBackground: true});
    await p.screenshot({path: out + '/previews/' + n + '.png', fullPage: true});
    await p.close();
  }
  await b.close();
})();
"""


def render(pages):
    """Print each HTML page to PDF with the pre-installed Node Playwright."""
    (OUT / "previews").mkdir(parents=True, exist_ok=True)
    for name, html in pages:
        (OUT / f".{name}.html").write_text(html, encoding="utf-8")
    subprocess.run(["node", "-e", RENDER_JS, str(OUT), *[n for n, _ in pages]], check=True)
    for name, _ in pages:
        (OUT / f".{name}.html").unlink()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    letterhead()
    render([("Aura-Business-Cards", cards_html()), ("Aura-Compliment-Slip", slip_html()),
            ("Aura-Envelope-DL", envelope_html())])
    (OUT / "Aura-Email-Signatures.html").write_text(signatures_html(), encoding="utf-8")
    print("stationery built in", OUT)


if __name__ == "__main__":
    main()
