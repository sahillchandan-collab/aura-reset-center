import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = "/home/user/aura-reset-center/yontro-crm/Yontro_CRM.xlsx"
N = 500  # working rows per tab
LAST = N + 1

INK, ORANGE, BONE, GREY = "0B0B0B", "FF5A1F", "F4F2EE", "4A4744"
F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor=INK)
fhdr_fill = PatternFill("solid", fgColor=GREY)
calc_fill = PatternFill("solid", fgColor=BONE)
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
thin = Side(style="thin", color="D9D6D0")
box = Border(bottom=thin)

def fill(c):
    return PatternFill("solid", fgColor=c)

wb = Workbook()

# ---------------------------------------------------------------- Settings
st = wb.active
st.title = "Settings"
lists = {
    "A": ("Lead stages", ["New", "Contacted", "Replied", "Meeting booked", "Meeting held",
                          "Proposal sent", "Contract signed + paid", "Closed – No response", "Lost"]),
    "B": ("Inbound sources", ["Website form", "WhatsApp", "Instagram DM", "Referral", "Walk-in", "Phone call", "Other"]),
    "C": ("Outbound sources", ["Apollo", "LinkedIn", "Google Maps", "Instagram research", "Event / networking", "Other"]),
    "D": ("Industries", ["Hospitality / F&B", "Retail", "Real estate", "Entertainment", "Fitness / wellness",
                         "Beauty / salon", "Clinic / healthcare", "Education", "Automotive", "Other"]),
    "E": ("Team", ["Usama", "Lianna", "Azeem", "Owner"]),
    "I": ("Payment methods", ["Bank transfer", "Card", "Cheque", "Cash", "Payment link"]),
    "J": ("Client decision", ["Pending", "Approved", "Backed out"]),
    "K": ("Back-out reasons", ["Video style", "Price", "Timing", "Changed mind", "Business issue", "Other"]),
    "L": ("Refund status", ["Requested", "Approved", "Paid"]),
    "M": ("Shoot types", ["Free video", "Plan shoot", "Add-on", "Event coverage"]),
    "N": ("Shoot status", ["Scheduled", "Shot", "Editing", "Sent for approval", "Delivered", "Posted"]),
    "O": ("Lost / closed reasons", ["No response (4 follow-ups)", "Price", "Not ready", "Went with competitor",
                                   "In-house team", "Not a fit", "Other"]),
    "P": ("Wants more of", ["Walk-in customers", "Enquiries and bookings", "Online sales",
                            "Followers and reach", "A better-looking brand"]),
}
rng = {}
for col, (title, items) in lists.items():
    st[f"{col}1"] = title
    for i, v in enumerate(items, start=2):
        st[f"{col}{i}"] = v
    rng[title] = f"Settings!${col}$2:${col}${len(items)+1}"
# packages table
st["F1"], st["G1"], st["H1"] = "Package", "Monthly fee (AED, ex VAT)", "Videos / month"
for i, (p, fee, v) in enumerate([("Launch", 5000, 10), ("Growth", 7500, 15), ("Scale", 10000, 20)], start=2):
    st[f"F{i}"], st[f"G{i}"], st[f"H{i}"] = p, fee, v
    st[f"G{i}"].number_format = "#,##0"
    st[f"G{i}"].font = Font(name=F, size=10, color="0000FF")
    st[f"H{i}"].font = Font(name=F, size=10, color="0000FF")
st["F6"] = "Source: yontro.ae Packages section. Prices exclude 5% VAT."
rng["Packages"] = "Settings!$F$2:$F$4"
st["Q1"] = "Shoot times"
for i in range(25):
    c = st[f"Q{i+2}"]
    c.value = dt.time(8 + i // 2, 30 * (i % 2))
    c.number_format = "hh:mm"
st["W1"] = "Times (all day)"
for i in range(48):
    c = st[f"W{i+2}"]
    c.value = dt.time(i // 2, 30 * (i % 2))
    c.number_format = "hh:mm"
TIME_ALL, TIME_WORK = "Settings!$W$2:$W$49", "Settings!$Q$2:$Q$26"
# rules
st["R1"], st["S1"] = "Rule", "Value"
rules = [
    ("Follow-up gap (days)", 3, "Your rule: follow up every 3 days after a lead goes quiet."),
    ("Max follow-ups before closing", 4, "Your rule: 4 follow-ups, each with new value, then stop outreach."),
    ("Decision window after free video (working days)", 3,
     "NOT DECIDED YET - placeholder of 3 working days. Change this one cell and every deadline updates."),
    ("Free video delivery promise (working days)", 5, "From the website: delivered within five working days."),
    ("Minimum term (months)", 3, "From the website: three-month minimum, then month to month."),
    ("Max revision rounds per video", 2, "From the website FAQ: two rounds of changes included."),
    ("Lead score for HOT", 70, "Lead qualification: 70+ = Hot. Book the meeting, send the proposal."),
    ("Lead score for WARM", 40, "Lead qualification: 40-69 = Warm. Keep following up. Under 40 = Cold."),
]
for i, (k, v, note) in enumerate(rules, start=2):
    st[f"R{i}"], st[f"S{i}"], st[f"T{i}"] = k, v, note
    st[f"S{i}"].font = Font(name=F, size=10, color="0000FF", bold=True)
st["S4"].fill = fill("FFFF00")
st["S4"].comment = Comment("Placeholder - the decision window has not been agreed yet.", "Yontro CRM")
GAP, MAXF, WIN, DELIV, TERM, REV = ("Settings!$S$2", "Settings!$S$3", "Settings!$S$4",
                                    "Settings!$S$5", "Settings!$S$6", "Settings!$S$7")
HOT, WARM = "Settings!$S$8", "Settings!$S$9"

# lead qualification: question / answer / points (edit points on Settings, formulas follow)
QUAL = [
    ("Budget AED 5k+/mo?", [("Yes", 25), ("Maybe", 10), ("No", 0)]),
    ("Decision-maker?", [("Yes", 15), ("Influencer", 8), ("No", 0)]),
    ("Someone on camera?", [("Yes", 15), ("Maybe", 8), ("No", 0)]),
    ("OK with 3-month term?", [("Yes", 10), ("Unsure", 5), ("No", 0)]),
    ("Start timeline", [("This month", 15), ("1–3 months", 8), ("Later", 2)]),
    ("Social media gap", [("Big gap", 10), ("Some", 5), ("Already strong", 2)]),
]
st["Y1"], st["Z1"], st["AA1"] = "Qualification question", "Answer", "Points"
qrng, r_ = {}, 2
for q, answers in QUAL:
    qrng[q] = f"Settings!$Z${r_}:$Z${r_ + len(answers) - 1}"
    for a, pts in answers:
        st[f"Y{r_}"], st[f"Z{r_}"], st[f"AA{r_}"] = q, a, pts
        st[f"AA{r_}"].font = Font(name=F, size=10, color="0000FF", bold=True)
        r_ += 1
st[f"Y{r_}"], st[f"Z{r_}"], st[f"AA{r_}"] = "Industry fit (automatic)", "Core industry", 10
st[f"Y{r_+1}"], st[f"Z{r_+1}"], st[f"AA{r_+1}"] = "Industry fit (automatic)", "Other industry", 5
for k in (r_, r_ + 1):
    st[f"AA{k}"].font = Font(name=F, size=10, color="0000FF", bold=True)
FIT_CORE, FIT_OTHER = f"Settings!$AA${r_}", f"Settings!$AA${r_+1}"
QY, QZ, QA = (f"Settings!$Y$2:$Y${r_+1}", f"Settings!$Z$2:$Z${r_+1}", f"Settings!$AA$2:$AA${r_+1}")
st["AC1"] = "Core industries"
for i, ind in enumerate(["Hospitality / F&B", "Retail", "Real estate", "Entertainment", "Fitness / wellness"], start=2):
    st[f"AC{i}"] = ind
CORE = "Settings!$AC$2:$AC$6"
for c in st[1]:
    if c.value:
        c.font, c.fill = hdr_font, hdr_fill
for row in st.iter_rows(min_row=2):
    for c in row:
        if c.font == Font():
            c.font = body
        elif c.font.name != F:
            c.font = Font(name=F, size=10, bold=c.font.bold, color=c.font.color)
for col in "ABCDEFGHIJKLMNOP":
    st.column_dimensions[col].width = 22
st.column_dimensions["R"].width = 44
st.column_dimensions["S"].width = 8
st.column_dimensions["T"].width = 80
st.column_dimensions["W"].width = 16
st.column_dimensions["Q"].width = 14
st.column_dimensions["Y"].width = 26
st.column_dimensions["Z"].width = 16
st.column_dimensions["AA"].width = 8
st.column_dimensions["AC"].width = 20
for c in st[1]:
    if c.value:
        c.font, c.fill = hdr_font, hdr_fill
st.freeze_panes = "A2"

# ---------------------------------------------------------------- helpers
def build_table(ws, cols, widths=None):
    """cols: list of (header, kind, extra) kind in input|calc; returns {header: letter}"""
    L = {}
    for i, (h, kind, *_) in enumerate(cols, start=1):
        L[h] = get_column_letter(i)
        c = ws.cell(row=1, column=i, value=h)
        c.font = hdr_font
        c.fill = fhdr_fill if kind == "calc" else hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L[h]].width = (widths or {}).get(h, 16)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{LAST}"
    return L

def style_rows(ws, cols, L, fmts):
    for r in range(2, LAST + 1):
        for h, kind, *_ in cols:
            c = ws[f"{L[h]}{r}"]
            c.font = body
            c.border = box
            if kind == "calc":
                c.fill = calc_fill
            if h in fmts:
                c.number_format = fmts[h]

def dv(ws, col, source, prompt=None):
    v = DataValidation(type="list", formula1=source, allow_blank=True)
    v.error = "Pick a value from the list (edit lists on the Settings tab)."
    v.errorTitle = "Not on the list"
    ws.add_data_validation(v)
    v.add(f"{col}2:{col}{LAST}")

def date_dv(ws, col, single=False):
    """Date validation: Google Sheets shows its calendar picker when the cell is double-clicked."""
    v = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)",
                       allow_blank=True)
    v.error = "Enter a date - double-click the cell for a calendar."
    v.errorTitle = "Not a date"
    ws.add_data_validation(v)
    if single:
        v.add(col)
    else:
        v.add(f"{col}2:{col}{LAST}")
        ws[f"{col}1"].value = f"{ws[f'{col}1'].value} 📅"

def time_dv(ws, col, source):
    """Dropdown of time slots. Warning only, so an exact time can still be typed."""
    v = DataValidation(type="list", formula1=source, allow_blank=True, errorStyle="warning")
    v.error = "That time isn't in the list - keep it if it's correct."
    v.errorTitle = "Check the time"
    ws.add_data_validation(v)
    v.add(f"{col}2:{col}{LAST}")

def int_dv(ws, col, lo, hi):
    v = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    v.error = f"Enter a whole number from {lo} to {hi}."
    ws.add_data_validation(v)
    v.add(f"{col}2:{col}{LAST}")

def cf(ws, ref, formula, color=None, font_color=None, bold_=False):
    ws.conditional_formatting.add(ref, FormulaRule(
        formula=[formula],
        fill=fill(color) if color else None,
        font=Font(name=F, color=font_color, bold=bold_) if (font_color or bold_) else None))

CLOSED = '"Contract signed + paid","Closed – No response","Lost"'

# ---------------------------------------------------------------- lead tabs
def lead_tab(name, inbound):
    ws = wb.create_sheet(name)
    cols = [("Lead ID", "input"),
            ("Date received" if inbound else "Date added", "input")]
    if inbound:
        cols += [("Time received", "input")]
    cols += [
            ("Source", "input"), ("Business name", "input"), ("Industry", "input"),
            ("Contact name", "input"), ("Role", "input"), ("Phone / WhatsApp", "input"),
            ("Email", "input"), ("Instagram", "input")]
    if not inbound:
        cols += [("LinkedIn", "input")]
    cols += [("Area", "input")]
    if inbound:
        cols += [("Wants more of", "input"), ("First response date", "input"), ("First response time", "input"),
                 ("Response time (hrs)", "calc")]
    else:
        cols += [("Apollo list / campaign", "input")]
    cols += [(q, "input") for q, _ in QUAL]
    cols += [("Lead score", "calc"), ("Rating", "calc"), ("Red flags", "calc")]
    cols += [("Owner", "input"), ("Stage", "input"), ("Last contact date", "input"),
             ("Follow-ups sent (0-4)", "input"), ("Next follow-up due", "calc"),
             ("Meeting date", "input"), ("Meeting time", "input"), ("Package discussed", "input"),
             ("Lost / closed reason", "input"), ("Notes", "input"),
             ("Days open", "calc"), ("Action flag", "calc")]
    widths = {"Business name": 24, "Email": 24, "Notes": 34, "Action flag": 30, "Stage": 22,
              "Apollo list / campaign": 22, "Lost / closed reason": 22, "Wants more of": 20,
              "Lead ID": 10, "Industry": 18, "Source": 15, "Red flags": 30, "Rating": 18,
              "Lead score": 9, **{q: 13 for q, _ in QUAL}}
    L = build_table(ws, cols, widths)
    d0 = "Date received" if inbound else "Date added"
    fmts = {d0: "ddd dd-mmm-yy", "Last contact date": "ddd dd-mmm-yy",
            "Next follow-up due": "ddd dd-mmm-yy", "Meeting date": "ddd dd-mmm-yy", "Meeting time": "hh:mm",
            "Days open": "0", "Phone / WhatsApp": "@"}
    if inbound:
        fmts.update({"Time received": "hh:mm", "First response date": "ddd dd-mmm-yy",
                     "First response time": "hh:mm", "Response time (hrs)": "0.0"})
    style_rows(ws, cols, L, fmts)

    for r in range(2, LAST + 1):
        b, stg, last, fu = (f"{L['Business name']}{r}", f"{L['Stage']}{r}",
                            f"{L['Last contact date']}{r}", f"{L['Follow-ups sent (0-4)']}{r}")
        ph, d = f"{L['Phone / WhatsApp']}{r}", f"{L[d0]}{r}"
        closed = f"OR({stg}=\"Contract signed + paid\",{stg}=\"Closed – No response\",{stg}=\"Lost\")"
        ws[f"{L['Next follow-up due']}{r}"] = (
            f'=IF(OR({b}="",{last}="",{closed},N({fu})>={MAXF}),"",{last}+{GAP})')
        ws[f"{L['Days open']}{r}"] = f'=IF(OR({b}="",{d}=""),"",INT(TODAY()-{d}))'
        refunded = (f'OR(COUNTIFS(Clients!$C$2:$C${LAST},{b},Clients!$S$2:$S${LAST},"Paid")>0,'
                    f'AND({ph}<>"",COUNTIFS(Clients!$E$2:$E${LAST},{ph},Clients!$S$2:$S${LAST},"Paid")>0))')
        nxt = f"{L['Next follow-up due']}{r}"
        qa = [f"{L[q]}{r}" for q, _ in QUAL]
        answered = f"COUNTA({qa[0]}:{qa[-1]})"
        ind = f"{L['Industry']}{r}"
        pts = "+".join(f'SUMIFS({QA},{QY},"{q}",{QZ},{c})' for (q, _), c in zip(QUAL, qa))
        fit = f'IF({ind}="",0,IF(COUNTIF({CORE},{ind})>0,{FIT_CORE},{FIT_OTHER}))'
        sc = f"{L['Lead score']}{r}"
        ws[sc] = f'=IF(OR({b}="",{answered}=0),"",{pts}+{fit})'
        ws[f"{L['Rating']}{r}"] = (
            f'=IF({b}="","",IF({answered}=0,"Not scored",IF({answered}<{len(QUAL)},'
            f'"Partly scored ("&{answered}&"/{len(QUAL)})",IF({sc}>={HOT},"Hot",IF({sc}>={WARM},"Warm","Cold")))))')
        bud, cam = f"{L['Budget AED 5k+/mo?']}{r}", f"{L['Someone on camera?']}{r}"
        ws[f"{L['Red flags']}{r}"] = (
            f'=IF({b}="","",IF(AND({bud}="No",{cam}="No"),"🚩 Budget  ·  🚩 No one on camera",'
            f'IF({bud}="No","🚩 Budget under AED 5k",IF({cam}="No","🚩 No one on camera",""))))')
        ws[f"{L['Action flag']}{r}"] = (
            f'=IF({b}="","",IF({refunded},"⚠ Refunded before – no free video",'
            f'IF({closed},"",IF({stg}="","Set a stage",'
            f'IF(AND({stg}="Proposal sent",{answered}<{len(QUAL)}),"Score this lead – proposal sent unscored",'
            f'IF(N({fu})>={MAXF},"Close – 4 follow-ups sent",'
            f'IF({nxt}="",IF({last}="","Make first contact",""),'
            f'IF(TODAY()>{nxt},"OVERDUE follow-up",IF(TODAY()={nxt},"Follow up today",""))))))))')
        if inbound:
            fr, ft, tr = (f"{L['First response date']}{r}", f"{L['First response time']}{r}",
                          f"{L['Time received']}{r}")
            ws[f"{L['Response time (hrs)']}{r}"] = (
                f'=IF(OR({d}="",{fr}=""),"",ROUND((({fr}+N({ft}))-({d}+N({tr})))*24,1))')

    dv(ws, L["Source"], rng["Inbound sources" if inbound else "Outbound sources"])
    dv(ws, L["Industry"], rng["Industries"])
    dv(ws, L["Owner"], rng["Team"])
    dv(ws, L["Stage"], rng["Lead stages"])
    dv(ws, L["Package discussed"], rng["Packages"])
    dv(ws, L["Lost / closed reason"], rng["Lost / closed reasons"])
    if inbound:
        dv(ws, L["Wants more of"], rng["Wants more of"])
    int_dv(ws, L["Follow-ups sent (0-4)"], 0, 4)
    for q, _ in QUAL:
        dv(ws, L[q], qrng[q])
    for h in (d0, "Last contact date", "Meeting date") + (("First response date",) if inbound else ()):
        date_dv(ws, L[h])
    time_dv(ws, L["Meeting time"], TIME_WORK)
    if inbound:
        time_dv(ws, L["Time received"], TIME_ALL)
        time_dv(ws, L["First response time"], TIME_ALL)

    # formatting rules
    s = L["Stage"]
    full = f"A2:{get_column_letter(len(cols))}{LAST}"
    cf(ws, full, f'${s}2="Contract signed + paid"', "DFF3E3")
    cf(ws, full, f'OR(${s}2="Closed – No response",${s}2="Lost")', None, "9A958E")
    fl = f"{L['Action flag']}2:{L['Action flag']}{LAST}"
    cf(ws, fl, f'LEFT({L["Action flag"]}2,1)="⚠"', "FAD4D4", "9B1C1C", True)
    cf(ws, fl, f'{L["Action flag"]}2="OVERDUE follow-up"', "FF5A1F", "FFFFFF", True)
    cf(ws, fl, f'{L["Action flag"]}2="Follow up today"', "FFE3B3", None, True)
    cf(ws, fl, f'{L["Action flag"]}2="Close – 4 follow-ups sent"', "E4DCF2", "4B2E83", True)
    cf(ws, fl, f'OR({L["Action flag"]}2="Make first contact",{L["Action flag"]}2="Set a stage")', "E6F0FA")
    cf(ws, fl, f'LEFT({L["Action flag"]}2,5)="Score"', "FFE3B3", "7A4B00", True)
    rt = L["Rating"]
    rr = f"{rt}2:{rt}{LAST}"
    cf(ws, rr, f'{rt}2="Hot"', "FF5A1F", "FFFFFF", True)
    cf(ws, rr, f'{rt}2="Warm"', "FFE3B3", "7A4B00", True)
    cf(ws, rr, f'{rt}2="Cold"', "DCE7F2", "1F3A5F", True)
    cf(ws, rr, f'OR({rt}2="Not scored",LEFT({rt}2,6)="Partly")', None, "9A958E")
    rf = L["Red flags"]
    cf(ws, f"{rf}2:{rf}{LAST}", f'{rf}2<>""', "FAD4D4", "9B1C1C", True)

    # example row
    ex = dict(zip([c[0] for c in cols], [None] * len(cols)))
    if inbound:
        ex.update({"Lead ID": "IN-0001", d0: dt.date(2026, 10, 15), "Time received": dt.time(10, 0),
                   "Source": "Website form",
                   "Business name": "Pulse Fitness Club (EXAMPLE)", "Industry": "Fitness / wellness",
                   "Contact name": "Sara M.", "Role": "Owner", "Phone / WhatsApp": "+971 50 000 0000",
                   "Email": "sara@example.com", "Instagram": "@pulse.fitness.jlt", "Area": "JLT",
                   "Wants more of": "Enquiries and bookings", "First response date": dt.date(2026, 10, 15),
                   "First response time": dt.time(10, 30),
                   "Owner": "Usama", "Stage": "Proposal sent", "Last contact date": dt.date(2026, 10, 17),
                   "Follow-ups sent (0-4)": 1, "Meeting date": dt.date(2026, 10, 16),
                   "Meeting time": dt.time(15, 0),
                   "Package discussed": "Growth", "Notes": "Example row – overwrite or delete.",
                   "Budget AED 5k+/mo?": "Yes", "Decision-maker?": "Yes", "Someone on camera?": "Yes",
                   "OK with 3-month term?": "Yes", "Start timeline": "This month", "Social media gap": "Big gap"})
    else:
        ex.update({"Lead ID": "OUT-0001", d0: dt.date(2026, 10, 13), "Source": "Apollo",
                   "Business name": "Marina Bites Café (EXAMPLE)", "Industry": "Hospitality / F&B",
                   "Contact name": "Omar K.", "Role": "Marketing manager", "Phone / WhatsApp": "+971 55 000 0000",
                   "Email": "omar@example.com", "Instagram": "@marinabites", "LinkedIn": "linkedin.com/in/example",
                   "Area": "Dubai Marina", "Apollo list / campaign": "F&B – Marina – Oct 26", "Owner": "Usama",
                   "Stage": "Replied", "Last contact date": dt.date(2026, 10, 16), "Follow-ups sent (0-4)": 0,
                   "Package discussed": "Launch", "Notes": "Example row – overwrite or delete.",
                   "Budget AED 5k+/mo?": "Maybe", "Decision-maker?": "Influencer", "Someone on camera?": "No",
                   "OK with 3-month term?": "Unsure", "Start timeline": "1–3 months", "Social media gap": "Some"})
    for h, v in ex.items():
        if v is not None and dict(cols)[h] == "input":
            ws[f"{L[h]}2"] = v
    for h, _ in cols:
        ws[f"{L[h]}2"].font = Font(name=F, size=10, italic=True, color="6B665F")
    return ws, L

inb, LI = lead_tab("Inbound", True)
outb, LO = lead_tab("Outbound", False)

# ---------------------------------------------------------------- Clients
cl = wb.create_sheet("Clients")
ccols = [("Client ID", "input"), ("Lead ID", "input"), ("Business name", "input"), ("Contact name", "input"),
         ("Phone / WhatsApp", "input"), ("Email", "input"), ("Instagram", "input"), ("Package", "input"),
         ("Monthly fee (AED)", "calc"), ("Contract signed date", "input"), ("Amount paid (AED)", "input"),
         ("Payment method", "input"), ("Free video shoot date", "input"), ("Free video delivered", "input"),
         ("Decision deadline", "calc"), ("Client decision", "input"), ("Status", "calc"),
         ("Back-out reason", "input"), ("Refund status", "input"), ("Refund amount (AED)", "input"),
         ("Refund date", "input"), ("Plan start date", "input"), ("Minimum term ends", "calc"),
         ("Account manager", "input"), ("Videographer", "input"), ("Notes", "input")]
LC = build_table(cl, ccols, {"Business name": 24, "Status": 34, "Notes": 34, "Email": 22, "Client ID": 10})
assert LC["Business name"] == "C" and LC["Phone / WhatsApp"] == "E" and LC["Refund status"] == "S"
style_rows(cl, ccols, LC, {"Monthly fee (AED)": "#,##0", "Amount paid (AED)": "#,##0",
                           "Refund amount (AED)": "#,##0", "Contract signed date": "ddd dd-mmm-yy",
                           "Free video shoot date": "ddd dd-mmm-yy", "Free video delivered": "ddd dd-mmm-yy",
                           "Decision deadline": "ddd dd-mmm-yy", "Refund date": "ddd dd-mmm-yy",
                           "Plan start date": "ddd dd-mmm-yy", "Minimum term ends": "ddd dd-mmm-yy",
                           "Phone / WhatsApp": "@"})
for r in range(2, LAST + 1):
    g = lambda h: f"{LC[h]}{r}"
    cl[g("Monthly fee (AED)")] = (f'=IF({g("Package")}="","",IFERROR(INDEX(Settings!$G$2:$G$4,'
                                  f'MATCH({g("Package")},Settings!$F$2:$F$4,0)),""))')
    cl[g("Decision deadline")] = f'=IF({g("Free video delivered")}="","",WORKDAY({g("Free video delivered")},{WIN}))'
    dl = g("Decision deadline")
    cl[g("Status")] = (
        f'=IF({g("Business name")}="","",'
        f'IF({g("Client decision")}="Backed out",IF({g("Refund status")}="Paid","Refunded – closed",'
        f'IF({g("Back-out reason")}="","Refund – log reason first","Refund in progress")),'
        f'IF({g("Amount paid (AED)")}="","Awaiting payment",'
        f'IF({g("Free video delivered")}="",IF({g("Free video shoot date")}="","Schedule free video",'
        f'"Free video in production"),'
        f'IF({g("Client decision")}="Approved","Active",'
        f'IF(TODAY()>{dl},"Active – plan went ahead",'
        f'"Decision pending – "&NETWORKDAYS(TODAY(),{dl})&" working day(s) left"))))))')
    cl[g("Minimum term ends")] = f'=IF({g("Plan start date")}="","",EDATE({g("Plan start date")},{TERM}))'
dv(cl, LC["Package"], rng["Packages"])
dv(cl, LC["Payment method"], rng["Payment methods"])
dv(cl, LC["Client decision"], rng["Client decision"])
dv(cl, LC["Back-out reason"], rng["Back-out reasons"])
dv(cl, LC["Refund status"], rng["Refund status"])
dv(cl, LC["Account manager"], rng["Team"])
dv(cl, LC["Videographer"], rng["Team"])
for h in ("Contract signed date", "Free video shoot date", "Free video delivered", "Refund date", "Plan start date"):
    date_dv(cl, LC[h])
S_ = LC["Status"]
sref = f"{S_}2:{S_}{LAST}"
cf(cl, sref, f'LEFT({S_}2,6)="Active"', "DFF3E3", "1E6B3A", True)
cf(cl, sref, f'LEFT({S_}2,8)="Decision"', "FFE3B3", None, True)
cf(cl, sref, f'LEFT({S_}2,6)="Refund"', "FAD4D4", "9B1C1C", True)
cf(cl, sref, f'{S_}2="Awaiting payment"', "FF5A1F", "FFFFFF", True)
ex = {"Client ID": "CL-0001", "Lead ID": "IN-0001", "Business name": "Example Salon (EXAMPLE)",
      "Contact name": "Layla K.", "Phone / WhatsApp": "+971 52 000 0000", "Email": "layla@example.com",
      "Instagram": "@examplesalon", "Package": "Growth", "Contract signed date": dt.date(2026, 10, 12),
      "Amount paid (AED)": 7875, "Payment method": "Bank transfer", "Free video shoot date": dt.date(2026, 10, 14),
      "Free video delivered": dt.date(2026, 10, 19), "Client decision": "Pending",
      "Account manager": "Lianna", "Videographer": "Azeem",
      "Notes": "Example row – overwrite or delete. Paid = AED 7,500 + 5% VAT."}
for h, v in ex.items():
    cl[f"{LC[h]}2"] = v
for h, _ in ccols:
    cl[f"{LC[h]}2"].font = Font(name=F, size=10, italic=True, color="6B665F")

# ---------------------------------------------------------------- Shoots
sh = wb.create_sheet("Shoots")
scols = [("Shoot ID", "input"), ("Client ID", "input"), ("Business name", "calc"), ("Shoot type", "input"),
         ("Shoot date", "input"), ("Time", "input"), ("Location", "input"), ("Videographer", "input"),
         ("Status", "input"), ("Delivered date", "input"), ("Revision rounds used", "input"),
         ("Video link", "input"), ("Notes", "input"), ("Flag", "calc")]
LS = build_table(sh, scols, {"Business name": 24, "Location": 22, "Video link": 26, "Notes": 30, "Flag": 30})
sh.freeze_panes = "D2"
style_rows(sh, scols, LS, {"Shoot date": "ddd dd-mmm-yy", "Delivered date": "ddd dd-mmm-yy", "Time": "hh:mm"})
for r in range(2, LAST + 1):
    g = lambda h: f"{LS[h]}{r}"
    sh[g("Business name")] = (f'=IF({g("Client ID")}="","",IFERROR(INDEX(Clients!$C$2:$C${LAST},'
                              f'MATCH({g("Client ID")},Clients!$A$2:$A${LAST},0)),"Client ID not found"))')
    sh[g("Flag")] = (
        f'=IF({g("Shoot ID")}="","",IF(N({g("Revision rounds used")})>{REV},"Over revision limit",'
        f'IF(AND({g("Shoot type")}="Free video",{g("Delivered date")}="",{g("Shoot date")}<>"",'
        f'TODAY()>WORKDAY({g("Shoot date")},{DELIV})),"LATE – free video promised in 5 working days",'
        f'IF(AND({g("Shoot date")}<>"",{g("Shoot date")}=TODAY()),"Shooting today",""))))')
dv(sh, LS["Shoot type"], rng["Shoot types"])
dv(sh, LS["Videographer"], rng["Team"])
dv(sh, LS["Status"], rng["Shoot status"])
int_dv(sh, LS["Revision rounds used"], 0, 9)
date_dv(sh, LS["Shoot date"])
date_dv(sh, LS["Delivered date"])
time_dv(sh, LS["Time"], TIME_WORK)
fl = f"{LS['Flag']}2:{LS['Flag']}{LAST}"
cf(sh, fl, f'LEFT({LS["Flag"]}2,4)="LATE"', "FF5A1F", "FFFFFF", True)
cf(sh, fl, f'{LS["Flag"]}2="Over revision limit"', "FAD4D4", "9B1C1C", True)
cf(sh, fl, f'{LS["Flag"]}2="Shooting today"', "FFE3B3", None, True)
ex = {"Shoot ID": "SH-0001", "Client ID": "CL-0001", "Shoot type": "Free video", "Shoot date": dt.date(2026, 10, 14),
      "Time": dt.time(10, 0), "Location": "Client salon, Jumeirah", "Videographer": "Azeem", "Status": "Delivered",
      "Delivered date": dt.date(2026, 10, 19), "Revision rounds used": 1, "Notes": "Example row – overwrite or delete."}
for h, v in ex.items():
    sh[f"{LS[h]}2"] = v
for h, _ in scols:
    sh[f"{LS[h]}2"].font = Font(name=F, size=10, italic=True, color="6B665F")

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 0)
db.sheet_view.showGridLines = False
db.column_dimensions["A"].width = 3
db.column_dimensions["B"].width = 40
for c in "CDE":
    db.column_dimensions[c].width = 16
db.column_dimensions["F"].width = 4
db.column_dimensions["G"].width = 30
for c in "HIJ":
    db.column_dimensions[c].width = 14
db["B1"] = "YONTRO · CRM DASHBOARD"
db["B1"].font = Font(name=F, size=18, bold=True, color=INK)
db["B2"] = "Updates automatically from the Inbound, Outbound, Clients and Shoots tabs. Only the yellow cell is yours to change."
db["B2"].font = Font(name=F, size=10, italic=True, color="6B665F")
db["B3"], db["D3"] = "Day to view 📅 (double-click)", "Leave blank to always show today"
db["B4"], db["C4"] = "Showing", "=IF(C3=\"\",TODAY(),C3)"
db["B5"], db["C5"] = "Week starts (Monday)", "=C4-WEEKDAY(C4,2)+1"
db["B6"], db["C6"] = "Month starts", "=DATE(YEAR(C4),MONTH(C4),1)"
for r in (3, 4, 5, 6):
    db[f"B{r}"].font = body
    db[f"C{r}"].font = bold
    db[f"C{r}"].number_format = "ddd dd-mmm-yy"
db["C3"].fill = fill("FFFF00")
db["C3"].font = Font(name=F, size=10, bold=True, color="0000FF")
db["C3"].comment = Comment("Double-click for a calendar. Clear the cell to go back to today.", "Yontro CRM")
db["D3"].font = Font(name=F, size=9, italic=True, color="6B665F")
date_dv(db, "C3", single=True)

IB, OB = f"Inbound!$B$2:$B${LAST}", f"Outbound!$B$2:$B${LAST}"
IS, OS = f"Inbound!${LI['Stage']}$2:${LI['Stage']}${LAST}", f"Outbound!${LO['Stage']}$2:${LO['Stage']}${LAST}"
IF_, OF_ = (f"Inbound!${LI['Action flag']}$2:${LI['Action flag']}${LAST}",
            f"Outbound!${LO['Action flag']}$2:${LO['Action flag']}${LAST}")
IM, OM = (f"Inbound!${LI['Meeting date']}$2:${LI['Meeting date']}${LAST}",
          f"Outbound!${LO['Meeting date']}$2:${LO['Meeting date']}${LAST}")
IO, OO = f"Inbound!${LI['Owner']}$2:${LI['Owner']}${LAST}", f"Outbound!${LO['Owner']}$2:${LO['Owner']}${LAST}"
ISrc, OSrc = f"Inbound!${LI['Source']}$2:${LI['Source']}${LAST}", f"Outbound!${LO['Source']}$2:${LO['Source']}${LAST}"
INm, ONm = (f"Inbound!${LI['Business name']}$2:${LI['Business name']}${LAST}",
            f"Outbound!${LO['Business name']}$2:${LO['Business name']}${LAST}")
CS = f"Clients!${LC['Status']}$2:${LC['Status']}${LAST}"
CD = f"Clients!${LC['Client decision']}$2:${LC['Client decision']}${LAST}"
CJ = f"Clients!${LC['Contract signed date']}$2:${LC['Contract signed date']}${LAST}"
CK = f"Clients!${LC['Amount paid (AED)']}$2:${LC['Amount paid (AED)']}${LAST}"
CI = f"Clients!${LC['Monthly fee (AED)']}$2:${LC['Monthly fee (AED)']}${LAST}"
CRs, CRa, CRd = (f"Clients!${LC['Refund status']}$2:${LC['Refund status']}${LAST}",
                 f"Clients!${LC['Refund amount (AED)']}$2:${LC['Refund amount (AED)']}${LAST}",
                 f"Clients!${LC['Refund date']}$2:${LC['Refund date']}${LAST}")
CN = f"Clients!${LC['Free video delivered']}$2:${LC['Free video delivered']}${LAST}"
RT = f"Inbound!${LI['Response time (hrs)']}$2:${LI['Response time (hrs)']}${LAST}"

def section(row, title, col="B", span=4):
    c = db[f"{col}{row}"]
    c.value = title
    c.font = Font(name=F, size=11, bold=True, color="FFFFFF")
    start = ord(col) - 64
    for i in range(start, start + span):
        db.cell(row=row, column=i).fill = fill(INK)

def kpi(row, label, formula, fmt="0", col="B", vcol="C", note=None):
    db[f"{col}{row}"] = label
    db[f"{col}{row}"].font = body
    db[f"{vcol}{row}"] = formula
    db[f"{vcol}{row}"].font = Font(name=F, size=12, bold=True, color=INK)
    db[f"{vcol}{row}"].number_format = fmt
    db[f"{vcol}{row}"].alignment = Alignment(horizontal="right")
    for c in (f"{col}{row}", f"{vcol}{row}"):
        db[c].border = box
    if note:
        db[f"{vcol}{row}"].comment = Comment(note, "Yontro CRM")

def section(row, title, col="B", span=4):
    c = db[f"{col}{row}"]
    c.value = title
    c.font = Font(name=F, size=11, bold=True, color="FFFFFF")
    start = ord(col) - 64
    for i in range(start, start + span):
        db.cell(row=row, column=i).fill = fill(INK)

def kpi(row, label, formula, fmt="0", col="B", vcol="C", note=None):
    db[f"{col}{row}"] = label
    db[f"{col}{row}"].font = body
    db[f"{vcol}{row}"] = formula
    db[f"{vcol}{row}"].font = Font(name=F, size=12, bold=True, color=INK)
    db[f"{vcol}{row}"].number_format = fmt
    db[f"{vcol}{row}"].alignment = Alignment(horizontal="right")
    for c in (f"{col}{row}", f"{vcol}{row}"):
        db[c].border = box
    if note:
        db[f"{vcol}{row}"].comment = Comment(note, "Yontro CRM")

ILc, OLc = (f"Inbound!${LI['Last contact date']}$2:${LI['Last contact date']}${LAST}",
            f"Outbound!${LO['Last contact date']}$2:${LO['Last contact date']}${LAST}")
INx, ONx = (f"Inbound!${LI['Next follow-up due']}$2:${LI['Next follow-up due']}${LAST}",
            f"Outbound!${LO['Next follow-up due']}$2:${LO['Next follow-up due']}${LAST}")
IFr = f"Inbound!${LI['First response date']}$2:${LI['First response date']}${LAST}"
CDl = f"Clients!${LC['Decision deadline']}$2:${LC['Decision deadline']}${LAST}"
SDt = f"Shoots!${LS['Shoot date']}$2:${LS['Shoot date']}${LAST}"
SNm = f"Shoots!${LS['Shoot ID']}$2:${LS['Shoot ID']}${LAST}"

def split_rows(start, rows, orange=()):
    db[f"C{start-1}"], db[f"D{start-1}"], db[f"E{start-1}"] = "Inbound", "Outbound", "Total"
    for c in "CDE":
        db[f"{c}{start-1}"].font = Font(name=F, size=9, bold=True, color="6B665F")
        db[f"{c}{start-1}"].alignment = Alignment(horizontal="right")
    for i, (lab, fi, fo) in enumerate(rows, start=start):
        kpi(i, lab, fi)
        db[f"D{i}"], db[f"E{i}"] = fo, f"=C{i}+D{i}"
        for c in (f"D{i}", f"E{i}"):
            db[c].font = Font(name=F, size=12, bold=True, color=INK)
            db[c].border = box
        if lab in orange:
            db[f"E{i}"].font = Font(name=F, size=12, bold=True, color=ORANGE)

# Daily
section(8, "DAILY  ·  the day picked above")
split_rows(10, [
    ("New leads", f'=COUNTIFS({IB},">="&$C$4,{IB},"<"&$C$4+1,{INm},"<>")',
     f'=COUNTIFS({OB},">="&$C$4,{OB},"<"&$C$4+1,{ONm},"<>")'),
    ("Leads contacted", f'=COUNTIFS({ILc},$C$4,{INm},"<>")', f'=COUNTIFS({OLc},$C$4,{ONm},"<>")'),
    ("Follow-ups due", f'=COUNTIFS({INx},$C$4)', f'=COUNTIFS({ONx},$C$4)'),
    ("Sales meetings", f'=COUNTIFS({IM},">="&$C$4,{IM},"<"&$C$4+1)', f'=COUNTIFS({OM},">="&$C$4,{OM},"<"&$C$4+1)'),
    ("Overdue follow-ups (live, as of today)", f'=COUNTIF({IF_},"OVERDUE follow-up")',
     f'=COUNTIF({OF_},"OVERDUE follow-up")'),
], orange=("Overdue follow-ups (live, as of today)",))
kpi(15, "Inbound leads still waiting for a first reply (live)", f'=COUNTIFS({INm},"<>",{IFr},"")',
    note="Target: reply within 1 hour.")
db["C15"].font = Font(name=F, size=12, bold=True, color=ORANGE)
kpi(16, "Shoots", f'=COUNTIFS({SDt},$C$4,{SNm},"<>")')
kpi(17, "Free videos delivered", f'=COUNTIFS({CN},$C$4)')
kpi(18, "Client decisions due", f'=COUNTIFS({CDl},$C$4,{CS},"Decision*")')
kpi(19, "Contracts signed + paid", f'=COUNTIFS({CJ},$C$4)')
kpi(20, "Cash collected (AED)", f'=SUMIFS({CK},{CJ},$C$4)', "#,##0")
kpi(21, "Refunds waiting to be paid (live)", f'=COUNTIF({CRs},"Requested")+COUNTIF({CRs},"Approved")')

# This week
section(23, "THIS WEEK  ·  Monday to Sunday of the day picked")
split_rows(25, [
    ("Leads added", f'=COUNTIFS({IB},">="&$C$5,{IB},"<"&$C$5+7,{INm},"<>")',
     f'=COUNTIFS({OB},">="&$C$5,{OB},"<"&$C$5+7,{ONm},"<>")'),
    ("Sales meetings", f'=COUNTIFS({IM},">="&$C$5,{IM},"<"&$C$5+7)', f'=COUNTIFS({OM},">="&$C$5,{OM},"<"&$C$5+7)'),
    ("Follow-ups due today or overdue (live)",
     f'=COUNTIF({IF_},"Follow up today")+COUNTIF({IF_},"OVERDUE follow-up")',
     f'=COUNTIF({OF_},"Follow up today")+COUNTIF({OF_},"OVERDUE follow-up")'),
    ("Leads to close (4 follow-ups sent)",
     f'=COUNTIF({IF_},"Close – 4 follow-ups sent")', f'=COUNTIF({OF_},"Close – 4 follow-ups sent")'),
], orange=("Follow-ups due today or overdue (live)",))
kpi(29, "Avg inbound response time (hrs, all time)", f'=IFERROR(AVERAGE({RT}),"-")', "0.0",
    note="Target: under 1 hour.")

# This month
section(31, "THIS MONTH  ·  of the day picked")
mend = '"<"&EDATE($C$6,1)'
kpi(32, "Contracts signed + paid", f'=COUNTIFS({CJ},">="&$C$6,{CJ},{mend})')
kpi(33, "Cash collected at signing (AED)", f'=SUMIFS({CK},{CJ},">="&$C$6,{CJ},{mend})', "#,##0")
kpi(34, "Refunds paid (count)", f'=COUNTIFS({CRs},"Paid",{CRd},">="&$C$6,{CRd},{mend})')
kpi(35, "Refunds paid (AED)", f'=SUMIFS({CRa},{CRs},"Paid",{CRd},">="&$C$6,{CRd},{mend})', "#,##0")

# Clients
section(37, "CLIENTS  ·  live")
kpi(38, "Active clients", f'=COUNTIF({CS},"Active*")')
kpi(39, "Monthly recurring revenue (AED, ex VAT)", f'=SUMIFS({CI},{CS},"Active*")', "#,##0")
kpi(40, "Decisions pending", f'=COUNTIF({CS},"Decision*")')
kpi(41, "Free videos to schedule / in production",
    f'=COUNTIF({CS},"Schedule free video")+COUNTIF({CS},"Free video in production")')
kpi(42, "Trial-to-active rate", f'=IFERROR(COUNTIF({CS},"Active*")/(COUNTIF({CS},"Active*")+COUNTIF({CD},"Backed out")),"-")',
    "0%", note="Of clients whose free-video decision is settled, the share who stayed.")
kpi(43, "Signed but payment not logged", f'=COUNTIF({CS},"Awaiting payment")')

# Lead quality
IR, OR_ = f"Inbound!${LI['Rating']}$2:${LI['Rating']}${LAST}", f"Outbound!${LO['Rating']}$2:${LO['Rating']}${LAST}"
IRF, ORF = (f"Inbound!${LI['Red flags']}$2:${LI['Red flags']}${LAST}",
            f"Outbound!${LO['Red flags']}$2:${LO['Red flags']}${LAST}")
ISc, OSc = (f"Inbound!${LI['Lead score']}$2:${LI['Lead score']}${LAST}",
            f"Outbound!${LO['Lead score']}$2:${LO['Lead score']}${LAST}")
section(45, "LEAD QUALITY  ·  open leads, live")


def open_cnt(R, Sx, crit):
    crit = f'"{crit}"'
    return (f'COUNTIFS({R},{crit})-COUNTIFS({R},{crit},{Sx},"Contract signed + paid")'
            f'-COUNTIFS({R},{crit},{Sx},"Closed – No response")-COUNTIFS({R},{crit},{Sx},"Lost")')


split_rows(47, [
    ("Hot (70+)", f'={open_cnt(IR, IS, "Hot")}', f'={open_cnt(OR_, OS, "Hot")}'),
    ("Warm (40–69)", f'={open_cnt(IR, IS, "Warm")}', f'={open_cnt(OR_, OS, "Warm")}'),
    ("Cold (under 40)", f'={open_cnt(IR, IS, "Cold")}', f'={open_cnt(OR_, OS, "Cold")}'),
    ("Not scored or partly scored", f'={open_cnt(IR, IS, "Not scored")}+{open_cnt(IR, IS, "Partly*")}',
     f'={open_cnt(OR_, OS, "Not scored")}+{open_cnt(OR_, OS, "Partly*")}'),
    ("With a 🚩 red flag", f'={open_cnt(IRF, IS, "🚩*")}', f'={open_cnt(ORF, OS, "🚩*")}'),
], orange=("Hot (70+)",))
db["B53"] = "Win rate by rating (all time)"
db["B53"].font = Font(name=F, size=9, bold=True, color="6B665F")
for i, band in enumerate(("Hot", "Warm", "Cold"), start=54):
    won = (f'COUNTIFS({IR},"{band}",{IS},"Contract signed + paid")'
           f'+COUNTIFS({OR_},"{band}",{OS},"Contract signed + paid")')
    kpi(i, f"{band} leads that signed", f'=IFERROR(({won})/(COUNTIF({IR},"{band}")+COUNTIF({OR_},"{band}")),"-")', "0%",
        note="Shows whether the score predicts who signs. Adjust points on Settings if it doesn't.")

# Pipeline by stage
section(8, "PIPELINE BY STAGE  ·  live", "G", 4)
db["H9"], db["I9"], db["J9"] = "Inbound", "Outbound", "Total"
for c in ("H9", "I9", "J9"):
    db[c].font = Font(name=F, size=9, bold=True, color="6B665F")
    db[c].alignment = Alignment(horizontal="right")
stages = lists["A"][1]
for i, stg in enumerate(stages, start=10):
    db[f"G{i}"] = f"=Settings!A{i-8}"
    db[f"H{i}"] = f"=COUNTIF({IS},G{i})"
    db[f"I{i}"] = f"=COUNTIF({OS},G{i})"
    db[f"J{i}"] = f"=H{i}+I{i}"
    for c in "GHIJ":
        db[f"{c}{i}"].font = bold if c == "J" else body
        db[f"{c}{i}"].border = box
r0 = 10 + len(stages)
db[f"G{r0}"] = "Total leads"
db[f"H{r0}"], db[f"I{r0}"], db[f"J{r0}"] = f"=SUM(H10:H{r0-1})", f"=SUM(I10:I{r0-1})", f"=SUM(J10:J{r0-1})"
for c in "GHIJ":
    db[f"{c}{r0}"].font = bold

# Open leads by owner
section(r0 + 2, "OPEN LEADS BY OWNER  ·  live", "G", 4)
db[f"H{r0+3}"], db[f"I{r0+3}"] = "Open leads", "Actions due"
for c in (f"H{r0+3}", f"I{r0+3}"):
    db[c].font = Font(name=F, size=9, bold=True, color="6B665F")
for j, person in enumerate(lists["E"][1]):
    r = r0 + 4 + j
    db[f"G{r}"] = f"=Settings!E{j+2}"
    openf = lambda O, S, Nm: (f'COUNTIFS({O},G{r},{Nm},"<>")-COUNTIFS({O},G{r},{S},"Contract signed + paid")'
                              f'-COUNTIFS({O},G{r},{S},"Closed – No response")-COUNTIFS({O},G{r},{S},"Lost")')
    db[f"H{r}"] = f"={openf(IO, IS, INm)}+{openf(OO, OS, ONm)}"
    db[f"I{r}"] = (f'=COUNTIFS({IO},G{r},{IF_},"OVERDUE follow-up")+COUNTIFS({IO},G{r},{IF_},"Follow up today")'
                   f'+COUNTIFS({OO},G{r},{OF_},"OVERDUE follow-up")+COUNTIFS({OO},G{r},{OF_},"Follow up today")')
    for c in "GHI":
        db[f"{c}{r}"].font = body
        db[f"{c}{r}"].border = box

# Leads by source
rs = r0 + 10
section(rs, "LEADS BY SOURCE (ALL TIME)", "G", 4)
db[f"H{rs+1}"], db[f"I{rs+1}"], db[f"J{rs+1}"] = "Leads", "Signed", "Avg score"
srcs = [(s, "in") for s in lists["B"][1]] + [(s, "out") for s in lists["C"][1]]
for j, (s, kind) in enumerate(srcs):
    r = rs + 2 + j
    src, stg = (ISrc, IS) if kind == "in" else (OSrc, OS)
    db[f"G{r}"] = f"{'In' if kind == 'in' else 'Out'} · {s}"
    db[f"H{r}"] = f'=COUNTIF({src},"{s}")'
    db[f"I{r}"] = f'=COUNTIFS({src},"{s}",{stg},"Contract signed + paid")'
    db[f"J{r}"] = f'=IFERROR(AVERAGEIFS({ISc if kind == "in" else OSc},{src},"{s}"),"-")'
    db[f"J{r}"].number_format = "0"
    for c in "GHIJ":
        db[f"{c}{r}"].font = body
        db[f"{c}{r}"].border = box

# ---------------------------------------------------------------- Qualify script
qs = wb.create_sheet("Qualify")
qs.sheet_view.showGridLines = False
qs.column_dimensions["A"].width = 3
qs.column_dimensions["B"].width = 24
qs.column_dimensions["C"].width = 62
qs.column_dimensions["D"].width = 44
qs["B1"] = "QUALIFYING A LEAD – first call or chat (about 5 minutes)"
qs["B1"].font = Font(name=F, size=16, bold=True, color=INK)
qs["B2"] = ("Ask naturally, not like a form. Fill the six dropdowns on the lead's row; the score, rating and "
            "red flags work themselves out. Red flags are warnings, not a 'no' – the lead stays open.")
qs["B2"].font = Font(name=F, size=10, italic=True, color="6B665F")
qs["B2"].alignment = Alignment(wrap_text=True)
qs.merge_cells("B2:D2")
qs.row_dimensions[2].height = 30
for j, h in enumerate(("Column", "What to ask", "Listen for"), start=2):
    c = qs.cell(row=4, column=j, value=h)
    c.font, c.fill = hdr_font, hdr_fill
script = [
    ("Budget AED 5k+/mo?", "“Our plans start at AED 5,000 a month for 10 videos, fully done for you. Is that in the range you had in mind?”",
     "Yes / Maybe (‘depends what’s included’) / No. 🚩 if No."),
    ("Decision-maker?", "“Apart from you, who else would be involved in deciding?”",
     "Yes = they sign. Influencer = they recommend to someone else."),
    ("Someone on camera?", "“Our videos work because a real person from the business is in them – a manager, chef, coach, broker. Who could that be for you?”",
     "A name = Yes. ‘Maybe, we’d need to see’ = Maybe. 🚩 if nobody."),
    ("OK with 3-month term?", "“We work in three-month plans: month one fixes the basics, two finds the story, three scales it. Does that work for you?”",
     "Yes / Unsure / No."),
    ("Start timeline", "“If this felt right, when would you want the first shoot?”",
     "This month / 1–3 months / Later."),
    ("Social media gap", "Look at their Instagram before the call: faces on camera? posting in the last 2 weeks? a clear offer in the bio?",
     "Big gap = 2–3 missing. Some = 1 missing. Already strong = none."),
    ("Industry fit", "Nothing to ask – it scores itself from the Industry column.",
     "Core: hospitality, retail, real estate, entertainment, fitness."),
]
for i, (col, ask, hear) in enumerate(script, start=5):
    for j, v in enumerate((col, ask, hear), start=2):
        c = qs.cell(row=i, column=j, value=v)
        c.font = bold if j == 2 else body
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = box
    qs.row_dimensions[i].height = 46
qs["B13"] = "Score bands"
qs["B13"].font = bold
qs["B14"], qs["C14"] = "Hot  70–100", "Book the meeting now; proposal straight after."
qs["B15"], qs["C15"] = "Warm  40–69", "Keep the follow-up rhythm going; send a relevant result."
qs["B16"], qs["C16"] = "Cold  under 40", "Low priority. Usama decides whether to keep going."
qs["B17"], qs["C17"] = "🚩 Red flags", "Budget under AED 5k, or nobody on camera. A warning to handle in the conversation – not a disqualification."
for k in range(14, 18):
    qs[f"B{k}"].font, qs[f"C{k}"].font = bold, body
    qs[f"C{k}"].alignment = Alignment(wrap_text=True)
qs["B14"].fill, qs["B15"].fill, qs["B16"].fill, qs["B17"].fill = fill(ORANGE), fill("FFE3B3"), fill("DCE7F2"), fill("FAD4D4")
qs["B19"] = "Points and the Hot/Warm cut-offs live on the Settings tab (blue numbers) – change them there and every score updates."
qs["B19"].font = Font(name=F, size=9, italic=True, color="6B665F")

# ---------------------------------------------------------------- Automations
au = wb.create_sheet("Automations")
au.sheet_view.showGridLines = False
heads = ["#", "When this happens…", "App", "…this happens automatically", "Writes to (tab → columns)", "Phase"]
data = [
    (1, "Website form submitted", "yontro.ae → Make.com", "Add row, Source = Website form, Stage = New; alert Usama on WhatsApp",
     "Inbound → Lead ID, Date received, Time received, Source, Business, Contact, Phone, Instagram, Wants more of", "1"),
    (2, "Calendly meeting booked", "Calendly → Make.com", "Find lead by email/phone; set Stage = Meeting booked, Meeting date and Meeting time; reminder 24h before",
     "Inbound / Outbound → Stage, Meeting date, Meeting time", "1"),
    (3, "Prospects saved to an Apollo list", "Apollo → Make.com", "Add rows, skip duplicates by email or phone; Stage = New",
     "Outbound → Lead ID, Date added, Source, Business…Apollo list", "1"),
    (4, "Prospect replies to an Apollo sequence", "Apollo → Make.com", "Set Stage = Replied; Last contact date = today",
     "Outbound → Stage, Last contact date", "1"),
    (5, "Stage set to Contract signed + paid", "Sheets → Make.com", "Create Clients row and a Free video row in Shoots; send Lianna the onboarding checklist",
     "Clients → Client ID…Amount paid; Shoots → Shoot ID, Client ID, Shoot type", "1"),
    (6, "Shoot date entered", "Sheets → Google Calendar", "Create event on Azeem's calendar with location; reminder to client the day before",
     "(calendar only)", "1"),
    (7, "Free video delivered", "Sheets → Make.com", "Decision window starts; Lianna nudged halfway; Usama alerted on deadline day",
     "(reads Clients → Free video delivered, Decision deadline)", "2"),
    (8, "WhatsApp message from a new number", "WhatsApp Business API (e.g. Wati) → Make.com",
     "Add row, Source = WhatsApp; auto-reply with Calendly meeting link", "Inbound → Date received, Source, Phone, Stage", "2"),
    (9, "Instagram DM or comment keyword", "ManyChat → Make.com", "Add row, Source = Instagram DM; reply with meeting link",
     "Inbound → Date received, Source, Instagram, Stage", "3"),
    (10, "Every morning 8:00", "Apps Script / Make.com", "Message each owner their 'Follow up today' and 'OVERDUE' rows",
     "(reads Action flag)", "2"),
]
for j, h in enumerate(heads, start=1):
    c = au.cell(row=1, column=j, value=h)
    c.font, c.fill = hdr_font, hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
for i, row in enumerate(data, start=2):
    for j, v in enumerate(row, start=1):
        c = au.cell(row=i, column=j, value=v)
        c.font, c.border = body, box
        c.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDEF", (4, 32, 26, 52, 52, 7)):
    au.column_dimensions[col].width = w
r = len(data) + 3
au[f"A{r}"] = "Rules that keep it clean"
au[f"A{r}"].font = bold
for k, t in enumerate([
    "Each column has ONE writer: an app or a person, never both. Grey columns are formulas – nobody types there.",
    "Every lead gets a Lead ID (IN-0001 / OUT-0001). Phone or email is the duplicate check.",
    "Phase 1 = website form, Calendly, Apollo, Clients, Calendar. Add WhatsApp and Instagram once phase 1 runs smoothly.",
], start=1):
    au[f"A{r+k}"] = f"• {t}"
    au[f"A{r+k}"].font = body

# ---------------------------------------------------------------- Start Here
sh0 = wb.create_sheet("Start Here", 0)
sh0.sheet_view.showGridLines = False
sh0.column_dimensions["A"].width = 3
sh0.column_dimensions["B"].width = 26
sh0.column_dimensions["C"].width = 100
lines = [
    ("title", "YONTRO CRM"),
    ("sub", "One place for every lead – inbound or outbound – from first message to active client."),
    ("gap", None),
    ("h", "HOW TO USE THIS SHEET"),
    ("kv", ("White cells", "You type here. Most have dropdowns – pick from the list.")),
    ("kv", ("Grey cells / grey headers", "Formulas. Don't type here – they calculate themselves.")),
    ("kv", ("Row 2 on each tab", "An EXAMPLE row (grey italic). Overwrite or delete it once you start.")),
    ("kv", ("Settings tab", "Edit dropdown lists, package prices and the rules (follow-up gap, decision window).")),
    ("kv", ("Dates", "Columns marked 📅 open a calendar: double-click the cell and pick the day. Time columns have a dropdown arrow. If a calendar ever doesn't open, run the 'Add date pickers' script (see below).")),
    ("gap", None),
    ("h", "THE TABS"),
    ("kv", ("Dashboard", "Daily, weekly and monthly numbers. Pick any day in the yellow cell (blank = today). Everything else is read only.")),
    ("kv", ("Inbound", "Leads who came to us: website, WhatsApp, Instagram, referrals, walk-ins.  Owner: Usama")),
    ("kv", ("Outbound", "Leads we went after: Apollo, LinkedIn, Google Maps.  Owner: Usama")),
    ("kv", ("Clients", "Everyone who signed + paid: free video, decision window, refunds, active plan.  Owner: Lianna")),
    ("kv", ("Shoots", "Every shoot – free videos and plan shoots.  Owner: Azeem")),
    ("kv", ("Qualify", "The 6 questions Usama asks on the first call, what to listen for, and what the score means.")),
    ("kv", ("Automations", "What connects to what (Make.com, Calendly, Apollo, Calendar). The setup checklist.")),
    ("gap", None),
    ("h", "THE PIPELINE"),
    ("kv", ("Selling (Inbound / Outbound)", "New → Contacted → Replied → Meeting booked → Meeting held → Proposal sent → Contract signed + paid")),
    ("kv", ("", "or → Closed – No response (after 4 follow-ups) / Lost")),
    ("kv", ("Trial (Clients + Shoots)", "Signed + paid → Free video shot → Delivered → Decision window")),
    ("kv", ("Outcome (Clients)", "Approves → Active   ·   Says nothing → Active (plan goes ahead)   ·   Backs out → Refund, keeps the video")),
    ("gap", None),
    ("h", "THE RULES WE AGREED"),
    ("kv", ("Free video", "Only for clients who have signed the contract. As many as needed for onboarded clients.")),
    ("kv", ("Payment", "Taken at signing.")),
    ("kv", ("Backing out", "Client can back out after the free video if not satisfied → refund. They keep the free video.")),
    ("kv", ("Silence after delivery", "Plan goes ahead – they have signed and paid.")),
    ("kv", ("Ghosting (before contract)", "4 follow-ups, 3 days apart, each adding new value. Then stop outreach → Closed – No response.")),
    ("kv", ("Repeat refunds", "A business that was refunded before is flagged ⚠ – no second free video.")),
    ("kv", ("Lead score", "6 questions + industry = score out of 100. Hot 70+, Warm 40–69, Cold under 40. Score every lead before a proposal.")),
    ("kv", ("Red flags", "🚩 Budget under AED 5k, 🚩 nobody on camera. Warnings only – the lead is NOT disqualified.")),
    ("kv", ("Decision window", "NOT DECIDED YET – set to 3 working days as a placeholder (Settings tab, yellow cell).")),
    ("gap", None),
    ("h", "EACH PERSON'S DAY"),
    ("kv", ("Usama", "Inbound + Outbound: filter Owner = Usama and Action flag. Clear OVERDUE first, then 'Follow up today'. Reply to new inbound within 1 hour.")),
    ("kv", ("Azeem", "Shoots: filter Videographer = Azeem. Update Status and Delivered date the day it happens.")),
    ("kv", ("Lianna", "Clients: work every 'Decision pending' row; onboard new signings; log refunds with a reason.")),
    ("kv", ("Owner", "Dashboard once a week.")),
    ("gap", None),
    ("h", "IF THE CALENDARS DON'T OPEN (one-time fix, 2 minutes)"),
    ("kv", ("1", "In the Google Sheet: Extensions → Apps Script.")),
    ("kv", ("2", "Delete what's there, paste the code from yontro-crm/date-pickers.gs, click Save.")),
    ("kv", ("3", "Choose addDatePickers in the toolbar and click Run (allow access the first time).")),
    ("kv", ("4", "Back in the sheet, double-click any 📅 cell – the calendar opens.")),
    ("gap", None),
    ("h", "AFTER A FOLLOW-UP"),
    ("kv", ("Update two cells", "Last contact date = today, and Follow-ups sent +1. The sheet works out the next due date.")),
]
r = 1
for kind, v in lines:
    if kind == "title":
        sh0[f"B{r}"] = v
        sh0[f"B{r}"].font = Font(name=F, size=20, bold=True, color=INK)
    elif kind == "sub":
        sh0[f"B{r}"] = v
        sh0[f"B{r}"].font = Font(name=F, size=11, italic=True, color="6B665F")
    elif kind == "h":
        sh0[f"B{r}"] = v
        sh0[f"B{r}"].font = Font(name=F, size=11, bold=True, color="FFFFFF")
        sh0[f"B{r}"].fill = fill(ORANGE)
        sh0[f"C{r}"].fill = fill(ORANGE)
    elif kind == "kv":
        sh0[f"B{r}"], sh0[f"C{r}"] = v
        sh0[f"B{r}"].font = bold
        sh0[f"C{r}"].font = body
        sh0[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        sh0[f"B{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        sh0[f"B{r}"].border = sh0[f"C{r}"].border = box
    r += 1
for row in sh0.iter_rows(min_col=2, max_col=2):
    if row[0].value == "Decision window":
        sh0.cell(row=row[0].row, column=3).fill = fill("FFFF00")

# tab colours / order
for ws, colr in ((sh0, ORANGE), (db, INK), (inb, "2E7D32"), (outb, "1565C0"), (cl, "6A1B9A"),
                 (sh, "EF6C00"), (au, "757575"), (st, "9E9E9E")):
    ws.sheet_properties.tabColor = colr
qs.sheet_properties.tabColor = "C62828"
wb._sheets = [sh0, db, inb, outb, cl, sh, qs, au, st]
wb.active = 0
wb.save(OUT)
print("saved", OUT)
